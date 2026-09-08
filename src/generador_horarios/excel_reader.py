from __future__ import annotations

from datetime import date, datetime, time, timedelta
import logging
from pathlib import Path

from openpyxl import load_workbook

from .models import MatchRecord
from .normalization import normalize_lookup_key


REQUIRED_HEADERS = [
    "Categoría",
    "Local",
    "Visitante",
    "Fecha",
    "Hora",
    "Campo de juego",
]


def parse_excel_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()

    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_excel_time(value: object) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        total_minutes = round(float(value) * 24 * 60)
        hours, minutes = divmod(total_minutes, 60)
        hours = hours % 24
        return time(hour=hours, minute=minutes)

    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            continue
    return None


def _as_string(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_fill_rgb(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    if not text or text == "00000000":
        return ""
    return text


def _cell_has_playoff_fill(cell: object) -> bool:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type != "solid":
        return False

    fg_color = getattr(fill, "fgColor", None)
    if fg_color is None:
        return False

    rgb = _normalize_fill_rgb(getattr(fg_color, "rgb", None))
    if rgb.endswith("FFFF00"):
        return True

    indexed = getattr(fg_color, "indexed", None)
    return indexed == 13


def _row_has_playoff_highlight(row: tuple[object, ...], header_positions: dict[str, int]) -> bool:
    highlighted_cells = 0
    checked_cells = 0
    for header in REQUIRED_HEADERS:
        position = header_positions.get(header)
        if position is None or position >= len(row):
            continue
        checked_cells += 1
        if _cell_has_playoff_fill(row[position]):
            highlighted_cells += 1

    return checked_cells > 0 and highlighted_cells >= min(checked_cells, 4)


def load_matches(excel_path: Path, logger: logging.Logger) -> list[MatchRecord]:
    workbook = load_workbook(excel_path, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]

        header_row_cells = next(worksheet.iter_rows(min_row=1, max_row=1), None)
        header_row = tuple(cell.value for cell in header_row_cells) if header_row_cells is not None else None
        if header_row is None:
            raise ValueError("El Excel no contiene cabecera.")

        expected_keys = {normalize_lookup_key(header): header for header in REQUIRED_HEADERS}
        header_positions: dict[str, int] = {}
        for index, value in enumerate(header_row):
            key = normalize_lookup_key(_as_string(value))
            if key in expected_keys:
                header_positions[expected_keys[key]] = index

        missing_headers = [header for header in REQUIRED_HEADERS if header not in header_positions]
        if missing_headers:
            raise ValueError(f"Faltan columnas obligatorias: {', '.join(missing_headers)}")

        records: list[MatchRecord] = []
        for row_index, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            values = {
                header: row[position].value if position < len(row) else None
                for header, position in header_positions.items()
            }

            categoria = _as_string(values["Categoría"])
            local = _as_string(values["Local"])
            visitante = _as_string(values["Visitante"])
            campo = _as_string(values["Campo de juego"])
            fecha = parse_excel_date(values["Fecha"])
            hora = parse_excel_time(values["Hora"])

            record = MatchRecord(
                row_index=row_index,
                categoria=categoria,
                local_original=local,
                visitante_original=visitante,
                fecha=fecha,
                hora=hora,
                campo_original=campo,
                is_playoff=_row_has_playoff_highlight(row, header_positions),
            )

            missing_fields: list[str] = []
            if not categoria:
                missing_fields.append("Categoría")
            if not local:
                missing_fields.append("Local")
            if not visitante:
                missing_fields.append("Visitante")
            if fecha is None:
                missing_fields.append("Fecha")
            if hora is None:
                missing_fields.append("Hora")
            if not campo:
                missing_fields.append("Campo de juego")

            if missing_fields:
                record.status = "skipped"
                record.skip_reason = f"Datos críticos ausentes o inválidos: {', '.join(missing_fields)}"
                logger.warning("Fila %s omitida: %s", row_index, record.skip_reason)

            records.append(record)

        return records
    finally:
        workbook.close()
